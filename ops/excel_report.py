"""
ops/excel_report.py - Daily Excel trade report with performance metrics.

Generates/updates an Excel file with:
  - Sheet "Trades":        Every trade (appended daily).
  - Sheet "Daily Summary": One row per day with aggregated metrics.
"""
from __future__ import annotations

import os
from datetime import datetime, timezone
from pathlib import Path

import numpy as np

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# ── Column definitions ────────────────────────────────────────────
TRADE_HEADERS = [
    "Date", "Time (UTC)", "Ticket", "Strategy", "Direction",
    "Entry Price", "SL", "TP", "Exit Price",
    "Lots", "Risk %", "SL (pips)", "TP (pips)", "Spread (pips)",
    "PnL ($)", "Duration", "Close Reason", "News Status",
    "Signal Reason", "Equity", "Balance", "DD %",
]

SUMMARY_HEADERS = [
    "Date", "Trades", "Wins", "Losses", "Win Rate %",
    "Gross Profit", "Gross Loss", "Net PnL",
    "Profit Factor", "Expectancy",
    "Best Trade", "Worst Trade",
    "Max DD %", "Ending Balance",
]

# ── Styling constants ─────────────────────────────────────────────
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
_WIN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_LOSS_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_MONEY_FMT = '#,##0.00'
_PCT_FMT = '0.00"%"'


class ExcelReporter:
    """Writes daily trade data and metrics to an Excel workbook."""

    def __init__(self, cfg: dict):
        log_dir = Path(cfg["logging"]["log_dir"])
        log_dir.mkdir(parents=True, exist_ok=True)
        self.filepath = log_dir / "trading_report.xlsx"
        self._day_trades: list[dict] = []

    # ── Public API ────────────────────────────────────────────────

    def record_trade(self, trade_data: dict):
        """Buffer a completed trade for end-of-day export."""
        self._day_trades.append(trade_data)

    def export_daily(self, balance: float, peak_balance: float):
        """Write buffered trades + daily summary to Excel. Call at EOD."""
        if not OPENPYXL_AVAILABLE:
            return
        if not self._day_trades:
            return

        wb = self._load_or_create()
        self._append_trades(wb)
        self._append_summary(wb, balance, peak_balance)
        self._auto_column_widths(wb)
        wb.save(str(self.filepath))
        self._day_trades.clear()

    def has_pending_trades(self) -> bool:
        return len(self._day_trades) > 0

    # ── Workbook management ───────────────────────────────────────

    def _load_or_create(self) -> Workbook:
        if self.filepath.exists():
            return load_workbook(str(self.filepath))
        wb = Workbook()
        # --- Trades sheet ---
        ws_trades = wb.active
        ws_trades.title = "Trades"
        self._write_header_row(ws_trades, TRADE_HEADERS)
        # --- Daily Summary sheet ---
        ws_summary = wb.create_sheet("Daily Summary")
        self._write_header_row(ws_summary, SUMMARY_HEADERS)
        return wb

    # ── Trades sheet ──────────────────────────────────────────────

    def _append_trades(self, wb: Workbook):
        ws = wb["Trades"]
        for t in self._day_trades:
            row = [
                t.get("date", ""),
                t.get("time", ""),
                t.get("ticket", ""),
                t.get("strategy", ""),
                t.get("direction", ""),
                t.get("entry", ""),
                t.get("sl", ""),
                t.get("tp", ""),
                t.get("exit_price", ""),
                t.get("lots", ""),
                t.get("risk_pct", ""),
                t.get("sl_pips", ""),
                t.get("tp_pips", ""),
                t.get("spread", ""),
                t.get("pnl", ""),
                t.get("duration", ""),
                t.get("close_reason", ""),
                t.get("news_status", ""),
                t.get("signal_reason", ""),
                t.get("equity", ""),
                t.get("balance", ""),
                t.get("dd_pct", ""),
            ]
            row_idx = ws.max_row + 1
            ws.append(row)

            # Color-code PnL rows
            pnl = t.get("pnl", 0)
            if isinstance(pnl, (int, float)):
                fill = _WIN_FILL if pnl > 0 else _LOSS_FILL if pnl < 0 else None
                if fill:
                    for col in range(1, len(TRADE_HEADERS) + 1):
                        ws.cell(row=row_idx, column=col).fill = fill

            # Apply borders and number formats
            for col in range(1, len(TRADE_HEADERS) + 1):
                cell = ws.cell(row=row_idx, column=col)
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(horizontal="center")

            # Money format for price/PnL columns
            for col_idx in [6, 7, 8, 9, 15, 20, 21]:  # Entry, SL, TP, Exit, PnL, Equity, Balance
                cell = ws.cell(row=row_idx, column=col_idx)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = _MONEY_FMT

    # ── Daily Summary sheet ───────────────────────────────────────

    def _append_summary(self, wb: Workbook, balance: float, peak_balance: float):
        ws = wb["Daily Summary"]

        pnls = [t["pnl"] for t in self._day_trades if isinstance(t.get("pnl"), (int, float))]
        total = len(pnls)
        if total == 0:
            return

        wins = sum(1 for p in pnls if p > 0)
        losses = total - wins
        win_rate = (wins / total * 100) if total > 0 else 0
        gross_profit = sum(p for p in pnls if p > 0)
        gross_loss = abs(sum(p for p in pnls if p <= 0))
        net_pnl = sum(pnls)
        profit_factor = (gross_profit / gross_loss) if gross_loss > 0 else float("inf")
        expectancy = np.mean(pnls) if pnls else 0
        best = max(pnls) if pnls else 0
        worst = min(pnls) if pnls else 0
        dd_pct = ((peak_balance - balance) / peak_balance * 100) if peak_balance > 0 else 0
        dd_pct = max(0, dd_pct)

        today_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")

        row = [
            today_str, total, wins, losses, round(win_rate, 2),
            round(gross_profit, 2), round(-gross_loss, 2), round(net_pnl, 2),
            round(profit_factor, 2) if profit_factor != float("inf") else "N/A",
            round(expectancy, 2),
            round(best, 2), round(worst, 2),
            round(dd_pct, 2), round(balance, 2),
        ]
        row_idx = ws.max_row + 1
        ws.append(row)

        # Styling
        net_fill = _WIN_FILL if net_pnl > 0 else _LOSS_FILL if net_pnl < 0 else None
        for col in range(1, len(SUMMARY_HEADERS) + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
            if net_fill:
                cell.fill = net_fill

        # Money format
        for col_idx in [6, 7, 8, 10, 11, 12, 14]:
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(cell.value, (int, float)):
                cell.number_format = _MONEY_FMT

    # ── Helpers ───────────────────────────────────────────────────

    @staticmethod
    def _write_header_row(ws, headers: list[str]):
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN
            cell.border = _THIN_BORDER
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    @staticmethod
    def _auto_column_widths(wb: Workbook):
        for ws in wb.worksheets:
            for col_cells in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col_cells[0].column)
                for cell in col_cells:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = min(max_len + 3, 30)
